"""Gera docs/Apresentacao_Atividade_Nota2.pptx (python-pptx), 12 slides, ~8 min.

Os números vêm de docs/evidencias/relatorio.json.
"""
import io
import json
import re
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

DOCS = Path(__file__).resolve().parent
EVID = DOCS / "evidencias"
SAIDA = DOCS / "Apresentacao_Atividade_Nota2.pptx"

NECESSARIOS = [EVID / "relatorio.json", EVID / "relatorio-html.png", DOCS / "diagrama_arquitetura.png",
               DOCS / "grafico_regioes.png", DOCS / "grafico_ranking.png", DOCS / "grafico_status.png"]
faltando = [str(p) for p in NECESSARIOS if not p.exists()]
if faltando:
    sys.exit("Arquivos ausentes: " + ", ".join(faltando))

rel = json.loads((EVID / "relatorio.json").read_text(encoding="utf-8"))
# testes-pytest.txt = saída de `python rodar_testes.py` (um "N passed" por serviço)
_testes = [int(x) for x in re.findall(r"(\d+) passed", (EVID / "testes-pytest.txt").read_text(encoding="utf-8-sig"))]
if len(_testes) != 3:
    sys.exit("evidencias/testes-pytest.txt deve conter o resultado dos 3 serviços (rodar_testes.py)")
N_TESTES = sum(_testes)
ind = rel["indicadores"]
R, REG, RANK, CRIT, OUT, BASE = (ind["resumo"], ind["regioes"], ind["ranking_motoristas"],
                                 ind["atrasos_criticos"], ind["outliers_custo"], rel["base"])


def n(v, casas=2):
    s = f"{v:,.{casas}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


melhor, pior = REG[0], REG[-1]
lider = RANK[0]

AZUL = RGBColor(0x0B, 0x3D, 0x66)
AZUL_CLARO = RGBColor(0xDB, 0xE9, 0xF6)
LARANJA = RGBColor(0xC5, 0x5A, 0x11)
TEXTO = RGBColor(0x1A, 0x1A, 0x1A)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
AMARELO = RGBColor(0xFF, 0xFF, 0x00)
FONTE = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
LAYOUT_VAZIO = prs.slide_layouts[6]
TOTAL = 12


def caixa_texto(slide, x, y, w, h, fundo=None, ancora=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancora
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    if fundo is not None:
        tb.fill.solid()
        tb.fill.fore_color.rgb = fundo
    return tb


def escrever(par, texto, tam, negrito=False, cor=TEXTO, fonte=FONTE):
    r = par.add_run()
    r.text = texto
    r.font.size = Pt(tam)
    r.font.bold = negrito
    r.font.name = fonte
    r.font.color.rgb = cor
    return r


def marcador(par, nivel_recuo=0.3):
    pPr = par._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(nivel_recuo))))
    pPr.set("indent", str(-int(Inches(nivel_recuo))))
    for filho in list(pPr):
        if filho.tag in (qn("a:buNone"), qn("a:buChar")):
            pPr.remove(filho)
    bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
    pPr.append(bu)


def bullets(slide, x, y, w, h, linhas, tam=24):
    assert len(linhas) <= 6, "No máximo 6 marcadores por slide"
    assert tam >= 18
    tb = caixa_texto(slide, x, y, w, h)
    tf = tb.text_frame
    for i, linha in enumerate(linhas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        marcador(p)
        if isinstance(linha, tuple):  # (destaque, resto)
            escrever(p, linha[0], tam, True, AZUL)
            escrever(p, linha[1], tam)
        else:
            escrever(p, linha, tam)
    return tb


def novo_slide(titulo, notas, numero):
    s = prs.slides.add_slide(LAYOUT_VAZIO)
    barra = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.15))
    barra.fill.solid()
    barra.fill.fore_color.rgb = AZUL
    barra.line.fill.background()
    tb = caixa_texto(s, 0.5, 0.15, 12.3, 0.85, ancora=MSO_ANCHOR.MIDDLE)
    escrever(tb.text_frame.paragraphs[0], titulo, 34, True, BRANCO)
    rod = caixa_texto(s, 11.3, 7.0, 1.8, 0.4)
    rod.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    escrever(rod.text_frame.paragraphs[0], f"{numero} / {TOTAL}", 14, False, RGBColor(0x55, 0x55, 0x55))
    s.notes_slide.notes_text_frame.text = notas
    return s


def imagem(slide, origem, x, y, max_w, max_h):
    """Insere imagem ajustada à caixa (mantendo proporção), centralizada na caixa."""
    if isinstance(origem, (str, Path)):
        with Image.open(origem) as im:
            w, h = im.size
    else:
        w, h = Image.open(io.BytesIO(origem.getvalue())).size
    if not isinstance(origem, (str, Path)):
        origem.seek(0)
    esc = min(max_w / w, max_h / h)
    lw, lh = w * esc, h * esc
    px = x + (max_w - lw) / 2
    py = y + (max_h - lh) / 2
    return slide.shapes.add_picture(origem if not isinstance(origem, Path) else str(origem),
                                    Inches(px), Inches(py), Inches(lw), Inches(lh))


def cartao(slide, x, y, w, h, grande, legenda, cor=AZUL):
    f = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    f.fill.solid()
    f.fill.fore_color.rgb = AZUL_CLARO
    f.line.color.rgb = cor
    f.line.width = Pt(2)
    tf = f.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    escrever(p, grande, 30, True, cor)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    escrever(p2, legenda, 18, False, TEXTO)
    return f


def tabela(slide, x, y, w, linhas, larguras, tam=18, altura_linha=0.6):
    t = slide.shapes.add_table(len(linhas), len(linhas[0]), Inches(x), Inches(y), Inches(w),
                               Inches(altura_linha * len(linhas))).table
    for j, lw in enumerate(larguras):
        t.columns[j].width = Inches(lw)
    for i, linha in enumerate(linhas):
        for j, txt in enumerate(linha):
            c = t.cell(i, j)
            c.text = ""
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]
            escrever(p, txt, tam, i == 0, BRANCO if i == 0 else TEXTO)
            c.fill.solid()
            c.fill.fore_color.rgb = AZUL if i == 0 else (AZUL_CLARO if i % 2 else BRANCO)
    return t


# 1 -------------------------------------------------------------------------
s = prs.slides.add_slide(LAYOUT_VAZIO)
fundo = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(4.3))
fundo.fill.solid()
fundo.fill.fore_color.rgb = AZUL
fundo.line.fill.background()
tb = caixa_texto(s, 0.7, 0.5, 11.9, 3.5, ancora=MSO_ANCHOR.MIDDLE)
escrever(tb.text_frame.paragraphs[0], "Sistema de Microsserviços para Gestão e Análise de Desempenho de "
         "Entregas em Logística Urbana", 40, True, BRANCO)
p = tb.text_frame.add_paragraph()
p.space_before = Pt(12)
escrever(p, "Atividade para Nota 2 – Microsserviços com Docker e Docker Compose", 22, False, AZUL_CLARO)
nomes = caixa_texto(s, 0.7, 4.7, 8.2, 1.7, fundo=AMARELO)
for i in range(1, 5):
    p = nomes.text_frame.paragraphs[0] if i == 1 else nomes.text_frame.add_paragraph()
    escrever(p, f"[NOME COMPLETO DO ALUNO {i}] – RA: [RA]", 18)
inst = caixa_texto(s, 9.2, 4.7, 3.6, 0.9, fundo=AMARELO)
escrever(inst.text_frame.paragraphs[0], "[NOME DA INSTITUIÇÃO]", 18)
p = inst.text_frame.add_paragraph()
escrever(p, "[CIDADE], 2026", 18)
s.notes_slide.notes_text_frame.text = (
    "[Tempo: 0:20] Boa tarde. Somos o grupo formado por [NOMES]. Vamos apresentar um sistema de três "
    "microsserviços que se comunicam por HTTP, rodando em Docker Compose, para acompanhar o desempenho de "
    "entregas de uma transportadora urbana. A apresentação dura cerca de oito minutos.")

# 2 -------------------------------------------------------------------------
s = novo_slide("Cenário e problema", (
    "[Tempo: 0:40] O cenário é uma transportadora urbana com cinco regiões, frota de motos, vans e caminhões e "
    "uma equipe de motoristas. A gestão precisa saber se cumpre prazos, quanto custa cada quilômetro e onde "
    "estão os gargalos. Sem consolidação, as decisões são baseadas em percepção. Os dados são sintéticos, "
    "gerados de forma determinística, para que os resultados sejam reproduzíveis."), 2)
bullets(s, 0.7, 1.5, 12.0, 5.3, [
    ("Cenário: ", "transportadora urbana com 5 regiões, motos, vans e caminhões"),
    ("Pergunta 1: ", "qual percentual das entregas chega no prazo?"),
    ("Pergunta 2: ", "onde a operação é mais lenta ou mais cara?"),
    ("Pergunta 3: ", "quais motoristas são mais pontuais?"),
    ("Pergunta 4: ", "há atrasos graves ou custos fora do padrão?"),
    ("Dados: ", "sintéticos e determinísticos (semente fixa)"),
])

# 3 -------------------------------------------------------------------------
s = novo_slide("Objetivos", (
    "[Tempo: 0:40] O objetivo geral é construir um sistema em microsserviços que gere um relatório consolidado. "
    "Os objetivos específicos seguem o enunciado: um serviço de dados com banco relacional, um de "
    "processamento que só conversa por HTTP, um de relatório que orquestra, e tudo containerizado com um "
    "Dockerfile por serviço e Docker Compose."), 3)
bullets(s, 0.7, 1.5, 12.0, 5.3, [
    ("Geral: ", "sistema em microsserviços que gera um relatório final de desempenho"),
    ("Dados: ", "PostgreSQL com endpoints de leitura e escrita"),
    ("Processamento: ", "consome os dados por HTTP e calcula indicadores"),
    ("Relatório: ", "orquestra os serviços e entrega JSON, CSV e HTML"),
    ("Infraestrutura: ", "um Dockerfile por serviço e Docker Compose"),
    ("Validação: ", f"{N_TESTES} testes automatizados e verificação de ponta a ponta"),
])

# 4 -------------------------------------------------------------------------
s = novo_slide("Arquitetura", (
    "[Tempo: 0:50] Este é o diagrama. O cliente fala só com o Serviço de Relatório, na porta 8003. Ele chama o "
    "Serviço de Processamento, porta 8002, que busca os dados no Serviço de Dados, porta 8001. O Relatório "
    "também consulta o endpoint de resumo do Dados. Apenas o Serviço de Dados acessa o PostgreSQL, que não tem "
    "porta publicada. Tudo roda na rede logistica-net do Compose."), 4)
imagem(s, DOCS / "diagrama_arquitetura.png", 0.4, 1.35, 8.2, 5.6)
bullets(s, 8.8, 1.7, 4.2, 5.0, [
    "Cliente acessa só o Relatório",
    "Serviços conversam por HTTP/JSON",
    "Só o Dados acessa o banco",
    "Rede interna no Compose",
], tam=20)

# 5 -------------------------------------------------------------------------
s = novo_slide("Tecnologias e justificativas", (
    "[Tempo: 0:40] Escolhemos o menor conjunto de ferramentas que cumpre os requisitos. Python com Flask pela "
    "simplicidade, Gunicorn para não usar o servidor de desenvolvimento, PostgreSQL pela integridade "
    "relacional, requests para as chamadas entre serviços, Docker e Compose para reprodutibilidade e pytest "
    "para testar os três serviços sem rede nem banco."), 5)
tabela(s, 0.7, 1.45, 11.9, [
    ["Tecnologia", "Por que foi escolhida"],
    ["Python + Flask", "API enxuta e fácil de explicar"],
    ["Gunicorn", "Servidor de produção nos contêineres"],
    ["PostgreSQL 16", "Relacional, com chaves e restrições"],
    ["requests", "Chamadas HTTP entre serviços"],
    ["Docker + Compose", "Ambiente idêntico com um comando"],
    ["pytest", "Testa os 3 serviços sem Docker, rede nem banco"],
], [3.9, 8.0], tam=20, altura_linha=0.72)

# 6 -------------------------------------------------------------------------
s = novo_slide("Serviço de Dados (porta 8001)", (
    "[Tempo: 0:45] O Serviço de Dados é o dono do banco. Há três tabelas: motoristas, veículos e entregas, com "
    "chaves estrangeiras e restrições. Na primeira subida ele carrega, com semente fixa 42, "
    f"{BASE['motoristas']} motoristas, {BASE['veiculos']} veículos e {BASE['entregas']} entregas, e não "
    "duplica nada nos reinícios. Há GET com filtros e paginação, POST e PATCH. Erros voltam em JSON: 400, 404, 405, 409 "
    "para placa duplicada, 500 e 503 quando o banco cai."), 6)
bullets(s, 0.7, 1.5, 12.0, 5.3, [
    ("Modelo: ", "motoristas, veículos e entregas com chaves estrangeiras"),
    ("Carga: ", f"{BASE['motoristas']} motoristas, {BASE['veiculos']} veículos, {BASE['entregas']} entregas"),
    ("Leitura: ", "GET com filtros (região, status, datas) e paginação"),
    ("Escrita: ", "POST em todas as entidades e PATCH em entregas"),
    ("Erros: ", "JSON com 400, 404, 405, 409 (placa duplicada), 500 e 503"),
    ("Idempotente: ", "reiniciar não duplica os dados"),
])

# 7 -------------------------------------------------------------------------
s = novo_slide("Serviço de Processamento (porta 8002)", (
    "[Tempo: 0:45] O Processamento não tem driver de banco: lê tudo por HTTP, paginando de mil em mil. A regra "
    "central é que o atraso é derivado: máximo entre zero e a diferença entre entrega real e prazo. Com isso "
    "calculamos percentual no prazo, atraso médio, custo por quilômetro, ranking de motoristas, atrasos "
    "críticos de 120 minutos ou mais e outliers de custo pela regra do IQR. As funções são puras e "
    "cobertas por testes automatizados. Ele espera 5 segundos pelo Dados, com 2 tentativas."), 7)
bullets(s, 0.7, 1.5, 12.0, 5.3, [
    ("Entrada: ", "somente HTTP, paginando o Serviço de Dados"),
    ("Atraso: ", "derivado, máx(0, entregue_em − prazo_previsto)"),
    ("Indicadores: ", "% no prazo, atraso médio e custo por km"),
    ("Ranking: ", "motoristas com pelo menos 5 entregas concluídas"),
    ("Críticos: ", "atraso de 120 min ou mais; outliers pela regra do IQR"),
    ("Qualidade e testes: ", f"{N_TESTES} testes pytest nos 3 serviços, sem Docker nem rede"),
])

# 8 -------------------------------------------------------------------------
s = novo_slide("Serviço de Relatório (porta 8003)", (
    "[Tempo: 0:35] O Relatório orquestra: chama o resumo no Dados, chama o endpoint completo no Processamento e "
    "monta o relatório consolidado, com observações automáticas. Ele entrega em JSON para integração, CSV "
    "para planilhas e HTML para leitura direta. Formato inválido devolve 400 e, se uma dependência falhar, "
    "devolve 502 identificando qual serviço caiu. Os timeouts são em cascata: 5 segundos do Processamento ao "
    "Dados, 30 do Relatório ao Processamento, e 60 no Gunicorn."), 8)
for i, (t, d) in enumerate([("JSON", "integração com outros sistemas"), ("CSV", "planilhas (formato longo)"),
                            ("HTML", "leitura direta no navegador")]):
    cartao(s, 0.7 + i * 4.1, 1.6, 3.8, 1.6, t, d)
bullets(s, 0.7, 3.7, 12.0, 3.0, [
    ("Fluxo: ", "Dados (/resumo) + Processamento (/metricas/completo) → relatório"),
    ("Entrada: ", "/relatorio?formato=json|csv|html"),
    ("Falhas: ", "400 para formato inválido; 502 se uma dependência cair"),
    ("Timeouts: ", "5 s (Processamento→Dados) < 30 s (Relatório→Processamento) < 60 s (Gunicorn)"),
])

# 9 -------------------------------------------------------------------------
s = novo_slide("Docker e Docker Compose", (
    "[Tempo: 0:50] Cada serviço tem o seu Dockerfile, com imagem python 3.12.12 slim, dependências fixadas e usuário "
    "sem privilégios. O Compose define a rede logistica-net, o volume do banco e quatro serviços. Os "
    "healthchecks usam urllib porque a imagem slim não tem curl, e depends_on com service_healthy garante a "
    "ordem: banco, Dados, Processamento e Relatório. O banco não publica porta no host."), 9)
bullets(s, 0.7, 1.5, 12.0, 5.3, [
    ("Dockerfiles: ", "3, um por serviço, base python:3.12.12-slim, usuário não root"),
    ("Compose: ", "projeto logistica, rede logistica-net, volume pgdata"),
    ("Healthchecks: ", "pg_isready no banco e /health nos serviços"),
    ("Ordem: ", "depends_on com service_healthy"),
    ("Exposição: ", "portas 8001, 8002 e 8003; o banco fica interno"),
    ("Subida: ", "docker compose up -d --build --wait"),
])

# 10 ------------------------------------------------------------------------
s = novo_slide("Demonstração", (
    "[Tempo: 0:55] Na demonstração ao vivo: subimos tudo com um comando, verificamos que os quatro contêineres "
    "estão saudáveis, consultamos o resumo do Dados, que devolve 400 entregas, e abrimos o relatório HTML no "
    "navegador. Depois paramos o Processamento para mostrar o 502 controlado e o retorno ao normal. À direita "
    "está a captura do relatório HTML gerado."), 10)
cod = caixa_texto(s, 0.5, 1.5, 6.4, 4.2, fundo=RGBColor(0x1E, 0x1E, 0x1E))
for i, linha in enumerate([
    "docker compose up -d --build --wait",
    "docker compose ps",
    "curl localhost:8001/resumo",
    "curl localhost:8002/metricas/completo",
    "curl \"localhost:8003/relatorio?formato=html\"",
]):
    p = cod.text_frame.paragraphs[0] if i == 0 else cod.text_frame.add_paragraph()
    p.space_after = Pt(12)
    escrever(p, linha, 18, False, RGBColor(0xE6, 0xE6, 0xE6), "Consolas")
bullets(s, 0.5, 5.8, 6.4, 1.2, ["Teste: parar o Processamento → 502"], tam=20)
with Image.open(EVID / "relatorio-html.png") as im:
    recorte = im.crop((0, 0, 1280, 1000)).convert("RGB")
buf = io.BytesIO()
recorte.save(buf, format="PNG")
imagem(s, buf, 7.2, 1.4, 5.7, 5.5)

# 11 ------------------------------------------------------------------------
s = novo_slide("Resultados", (
    f"[Tempo: 0:50] Dos números reais: {R['total']} entregas, das quais {R['finalizadas']} concluídas. "
    f"{n(R['pct_no_prazo'])} por cento chegaram no prazo; quando há atraso, ele é de "
    f"{n(R['atraso_medio_min'])} minutos em média; o custo é de {n(R['custo_por_km'])} reais por quilômetro. "
    f"{melhor['regiao']} é a região mais pontual ({n(melhor['pct_no_prazo'])}%) e {pior['regiao']} a menos "
    f"({n(pior['pct_no_prazo'])}%). {lider['nome']} lidera o ranking com {n(lider['pct_no_prazo'])}%. Foram "
    f"sinalizados {CRIT['total']} atrasos críticos e {OUT['total']} custos atípicos. Os testes de resiliência "
    "e idempotência passaram, e a recriação do ambiente reproduziu os mesmos números."), 11)
cartao(s, 0.5, 1.3, 3.0, 1.35, str(R["total"]), "entregas")
cartao(s, 3.6, 1.3, 3.0, 1.35, f"{n(R['pct_no_prazo'])}%", "no prazo")
cartao(s, 6.7, 1.3, 3.0, 1.35, f"{n(R['atraso_medio_min'])} min", "atraso médio")
cartao(s, 9.8, 1.3, 3.0, 1.35, f"R$ {n(R['custo_por_km'])}", "custo por km")
imagem(s, DOCS / "grafico_regioes.png", 0.4, 2.85, 6.3, 3.4)
imagem(s, DOCS / "grafico_ranking.png", 6.8, 2.85, 6.1, 3.4)
leg = caixa_texto(s, 0.4, 6.3, 12.5, 0.7)
escrever(leg.text_frame.paragraphs[0],
         f"Melhor região: {melhor['regiao']} ({n(melhor['pct_no_prazo'])}%) · Pior: {pior['regiao']} "
         f"({n(pior['pct_no_prazo'])}%) · Líder: {lider['nome']}", 18, True, AZUL)

# 12 ------------------------------------------------------------------------
s = novo_slide("Conclusão, limitações e próximos passos", (
    "[Tempo: 0:30] Concluímos que os objetivos foram atendidos: três microsserviços por HTTP, banco "
    "relacional, um Dockerfile por serviço e Compose. A limitação principal é que não há autenticação nem "
    "HTTPS; a senha do banco é didática, os dados são sintéticos, a comunicação é síncrona e há uma pequena "
    "duplicação deliberada de código entre serviços, para manter a independência. Como próximos passos: autenticação, fila de "
    "mensagens, cache e observabilidade. As referências estão no relatório: Fowler e Lewis, Newman e as "
    "documentações de Docker, Flask e PostgreSQL. Obrigado! Tempo total estimado: cerca de 8 minutos e 20 "
    "segundos, dentro da faixa de 5 a 10 minutos."), 12)
bullets(s, 0.7, 1.5, 12.0, 5.3, [
    ("Conclusão: ", "requisitos atendidos, com ambiente reproduzível em um comando"),
    ("Limitações: ", "sem autenticação nem HTTPS; senha didática; dados sintéticos"),
    ("Limitações: ", "comunicação síncrona; duplicação deliberada entre serviços"),
    ("Futuro: ", "autenticação, mensageria, cache, observabilidade e CI"),
    ("Referências: ", "Fowler e Lewis (2014); Newman (2021)"),
    ("Documentação: ", "Docker, Compose, Flask, PostgreSQL e Python"),
])

prs.save(SAIDA)
print(f"Gerado: {SAIDA}")


