"""Gera docs/Relatorio_Atividade_Nota2.docx (python-docx).

Todos os números do capítulo de Resultados são lidos de docs/evidencias/,
de modo que o texto nunca diverge dos dados capturados na execução real.
"""
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS = Path(__file__).resolve().parent
RAIZ = DOCS.parent
EVID = DOCS / "evidencias"
SAIDA = DOCS / "Relatorio_Atividade_Nota2.docx"

# ----------------------------------------------------------------------------
# Evidências (falha de forma explícita se algo estiver ausente)
# ----------------------------------------------------------------------------
OBRIGATORIOS = ["relatorio.json", "docker-compose-ps.txt", "resiliencia.txt", "testes-pytest.txt"]
FIGURAS = ["diagrama_arquitetura.png", "grafico_regioes.png", "grafico_status.png", "grafico_ranking.png"]
faltando = [f"evidencias/{n}" for n in OBRIGATORIOS if not (EVID / n).exists()]
faltando += [n for n in FIGURAS if not (DOCS / n).exists()]
if faltando:
    sys.exit("Arquivos ausentes (execute a pilha e os geradores anteriores): " + ", ".join(faltando))

rel = json.loads((EVID / "relatorio.json").read_text(encoding="utf-8"))
ind = rel["indicadores"]
R = ind["resumo"]
REGIOES = ind["regioes"]
RANK = ind["ranking_motoristas"]
CRIT = ind["atrasos_criticos"]
OUT = ind["outliers_custo"]
STATUS = ind["status"]
BASE = rel["base"]
ps_txt = (EVID / "docker-compose-ps.txt").read_text(encoding="utf-8-sig")
resil_txt = (EVID / "resiliencia.txt").read_text(encoding="utf-8-sig")
pytest_txt = (EVID / "testes-pytest.txt").read_text(encoding="utf-8-sig")
SCREENSHOT = EVID / "relatorio-html.png"


def n(v, casas=2):
    """Formata número no padrão pt-BR (vírgula decimal, ponto de milhar)."""
    s = f"{v:,.{casas}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def dh(iso):
    return datetime.fromisoformat(iso[:19]).strftime("%d/%m/%Y %H:%M")


def pct(v):
    return n(v) + "%"


# testes-pytest.txt é a saída de `python rodar_testes.py`: um "N passed" por serviço (dados, processamento, relatório)
m_testes = [int(x) for x in re.findall(r"(\d+) passed", pytest_txt)]
if len(m_testes) != 3 or re.search(r"\d+ (failed|error)", pytest_txt):
    sys.exit("evidencias/testes-pytest.txt deve conter o resultado aprovado dos 3 serviços (rodar_testes.py)")
T_DADOS, T_PROC, T_REL = m_testes
N_TESTES = sum(m_testes)
if "HTTP 502" not in resil_txt or "recuperacao: HTTP 200" not in resil_txt:
    sys.exit("evidencias/resiliencia.txt não contém o 502 esperado e a recuperação")

# Derivados (todos calculados a partir do relatorio.json)
melhor = REGIOES[0]
pior = REGIOES[-1]
menor_atraso = min(REGIOES, key=lambda r: r["atraso_medio_min"])
maior_atraso = max(REGIOES, key=lambda r: r["atraso_medio_min"])
mais_barata = min(REGIOES, key=lambda r: r["custo_por_km"])
mais_cara = max(REGIOES, key=lambda r: r["custo_por_km"])
lider = RANK[0]
ultimo = RANK[-1]
maior_atraso_motorista = max(RANK, key=lambda r: r["atraso_medio_min"])
pct_atrasadas = round(100 - R["pct_no_prazo"], 2)
nao_concluidas = R["total"] - R["finalizadas"]
crit_max = CRIT["itens"][0]
dif_regioes = round(melhor["pct_no_prazo"] - pior["pct_no_prazo"], 2)

# docker compose ps -> linhas simplificadas (nome, serviço, status, portas)
linhas_ps = [l for l in ps_txt.splitlines() if l.strip()]
cab = linhas_ps[0]
cols = {k: cab.index(k) for k in ["NAME", "IMAGE", "COMMAND", "SERVICE", "CREATED", "STATUS", "PORTS"]}
ordem_cols = sorted(cols.items(), key=lambda kv: kv[1])


def campo(linha, nome):
    i = [k for k, _ in ordem_cols].index(nome)
    ini = cols[nome]
    fim = ordem_cols[i + 1][1] if i + 1 < len(ordem_cols) else len(linha)
    return linha[ini:fim].strip()


TAB_PS = []
for l in linhas_ps[1:]:
    st = re.search(r"\((\w+)\)", campo(l, "STATUS"))
    portas = re.findall(r"0\.0\.0\.0:(\d+)->(\d+)", campo(l, "PORTS"))
    if portas:
        txt_portas = ", ".join(f"{a}→{b}" for a, b in portas[:1])
    else:
        txt_portas = "interna (5432), não publicada"
    TAB_PS.append([campo(l, "NAME"), campo(l, "SERVICE"), st.group(1) if st else campo(l, "STATUS"), txt_portas])
if len(TAB_PS) != 4 or any(r[2] != "healthy" for r in TAB_PS):
    sys.exit("docker-compose-ps.txt não mostra os 4 contêineres saudáveis")

DOCKERFILE = (RAIZ / "servico-dados" / "Dockerfile").read_text(encoding="utf-8").strip()

# ----------------------------------------------------------------------------
# Documento e estilos
# ----------------------------------------------------------------------------
doc = Document()
FONTE = "Arial"


def definir_fonte(estilo, tam=None, negrito=None, cor=None, nome=FONTE):
    f = estilo.font
    f.name = nome
    rpr = estilo.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), nome)
    if tam:
        f.size = Pt(tam)
    if negrito is not None:
        f.bold = negrito
    if cor is not None:
        f.color.rgb = cor


normal = doc.styles["Normal"]
definir_fonte(normal, 12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
doc.styles["Normal"].element.get_or_add_pPr()

for nome, tam, antes, depois in (("Heading 1", 14, 18, 8), ("Heading 2", 12, 12, 6), ("Heading 3", 12, 10, 4)):
    st = doc.styles[nome]
    definir_fonte(st, tam, True, RGBColor(0, 0, 0))
    st.font.italic = False
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.space_before = Pt(antes)
    st.paragraph_format.space_after = Pt(depois)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.line_spacing = 1.5
doc.styles["Heading 1"].font.all_caps = True
for nome in ("List Bullet", "List Number"):
    definir_fonte(doc.styles[nome], 12)

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin, sec.left_margin = Cm(3), Cm(3)
sec.bottom_margin, sec.right_margin = Cm(2), Cm(2)

cont = {"fig": 0, "tab": 0}
REGEX_MARCA = re.compile(r"(\*\*.+?\*\*|`.+?`|\[\[.+?\]\])")


def runs(p, texto, tam=None, negrito=False):
    """Escreve texto em p. **negrito**, `código`, [[marcador destacado em amarelo]]."""
    for parte in REGEX_MARCA.split(texto):
        if not parte:
            continue
        if parte.startswith("**"):
            r = p.add_run(parte[2:-2])
            r.bold = True
        elif parte.startswith("`"):
            r = p.add_run(parte[1:-1])
            r.font.name = "Courier New"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
            r.font.size = Pt((tam or 12) - 1)
            r.bold = negrito
            continue
        elif parte.startswith("[["):
            r = p.add_run("[" + parte[2:-2] + "]")
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            r.bold = negrito
        else:
            r = p.add_run(parte)
            r.bold = negrito or None
        if tam:
            r.font.size = Pt(tam)
    return p


def para(texto, recuo=True, alinhamento=None):
    p = doc.add_paragraph()
    if recuo:
        p.paragraph_format.first_line_indent = Cm(1.25)
    if alinhamento is not None:
        p.alignment = alinhamento
    return runs(p, texto)


def itens(lista, estilo="List Bullet"):
    for t in lista:
        p = doc.add_paragraph(style=estilo)
        runs(p, t)


def titulo(texto, nivel=1):
    return doc.add_heading(texto, level=nivel)


def sombrear(elemento_pr, cor):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    elemento_pr.append(shd)


def legenda(texto):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    runs(p, texto, tam=10, negrito=True)
    return p


def fonte_nota(texto="Fonte: elaborado pelos autores (2026)."):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(10)
    runs(p, texto, tam=10)


def figura(arquivo, largura_cm, texto, fonte=None):
    cont["fig"] += 1
    legenda(f"Figura {cont['fig']} – {texto}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(arquivo), width=Cm(largura_cm))
    fonte_nota(fonte) if fonte else fonte_nota()
    return cont["fig"]


def tabela(texto, cabecalho, linhas, larguras, alinh=None, fonte=None, tam=10):
    cont["tab"] += 1
    legenda(f"Tabela {cont['tab']} – {texto}")
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    alinh = alinh or ["l"] * len(cabecalho)
    mapa = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER}

    def preencher(celula, texto_cel, i, cab=False):
        celula.width = Cm(larguras[i])
        p = celula.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if cab else mapa[alinh[i]]
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.first_line_indent = Cm(0)
        runs(p, str(texto_cel), tam=tam, negrito=cab)

    for i, c in enumerate(cabecalho):
        preencher(t.rows[0].cells[i], c, i, cab=True)
        sombrear(t.rows[0].cells[i]._tc.get_or_add_tcPr(), "D9E2F3")
    trpr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    trpr.append(hdr)
    for linha in linhas:
        cells = t.add_row().cells
        for i, c in enumerate(linha):
            preencher(cells[i], c, i)
    for ri, row in enumerate(t.rows):
        cant = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant)
        if ri < len(t.rows) - 1 and len(t.rows) <= 14:  # mantém a tabela inteira na mesma página
            for c in row.cells:
                for cp in c.paragraphs:
                    cp.paragraph_format.keep_with_next = True
    fonte_nota(fonte) if fonte else fonte_nota()
    return cont["tab"]


def codigo(texto):
    for linha in texto.splitlines():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(0)
        sombrear(p._p.get_or_add_pPr(), "F2F2F2")
        r = p.add_run(linha if linha else " ")
        r.font.name = "Courier New"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        r.font.size = Pt(9)
    espaco = doc.add_paragraph()
    espaco.paragraph_format.line_spacing = 1.0


def campo_pagina(par):
    for tipo, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        r = par.add_run()
        if tipo:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tipo)
            r._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            r._r.append(it)
        r.font.size = Pt(10)


# ----------------------------------------------------------------------------
# CAPA (sem numeração de página)
# ----------------------------------------------------------------------------
def linha_capa(texto, tam=12, negrito=False, antes=0, depois=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(depois)
    runs(p, texto, tam=tam, negrito=negrito)
    return p


# Rótulo estrutural "Capa" (título de seção para navegação); oculto na impressão.
h_capa = titulo("Capa", 1)
for r in h_capa.runs:
    r.font.hidden = True
h_capa.paragraph_format.space_before = Pt(0)
h_capa.paragraph_format.space_after = Pt(0)

linha_capa("[[NOME DA INSTITUIÇÃO]]", 14, True)
linha_capa("[[CURSO / DISCIPLINA]]", 12, False, 4, 36)
for i in range(1, 5):
    linha_capa(f"[[NOME COMPLETO DO ALUNO {i}]] – RA: [[RA]]", 12, False, 0, 6)
linha_capa("Sistema de Microsserviços para Gestão e Análise de Desempenho de Entregas em Logística Urbana",
           18, True, 110, 18)
linha_capa("Atividade para Nota 2 – Arquitetura de microsserviços com Docker e Docker Compose", 12, False, 0, 0)
linha_capa("[[CIDADE]]", 12, False, 150, 0)
linha_capa("2026", 12, True, 0, 0)

# Nova seção (nova página) com rodapé numerado
nova = doc.add_section(WD_SECTION.NEW_PAGE)
nova.footer.is_linked_to_previous = False
fp = nova.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
campo_pagina(fp)

# ----------------------------------------------------------------------------
# 1 INTRODUÇÃO
# ----------------------------------------------------------------------------
titulo("1 Introdução", 1)
titulo("1.1 Cenário escolhido", 2)
para("O cenário escolhido é o de uma transportadora urbana de médio porte que realiza entregas de mercadorias "
     "em cinco regiões de uma grande cidade (Centro, Zona Norte, Zona Sul, Zona Leste e Zona Oeste), com uma frota "
     "composta por motos, vans e caminhões e uma equipe de motoristas. Cada entrega tem um prazo previsto, um "
     "veículo e um motorista responsáveis, uma distância percorrida e um custo associado. Ao final de cada ciclo, "
     "a gestão precisa saber se a operação está cumprindo os prazos, quanto custa cada quilômetro rodado e quais "
     "regiões e motoristas exigem atenção.")
para("Trata-se de um cenário não trivial porque envolve múltiplas entidades relacionadas (motoristas, veículos e "
     "entregas), regras de negócio derivadas (como o cálculo de atraso a partir de duas datas), agregações em "
     "diferentes níveis (geral, por região e por motorista) e a detecção de casos atípicos, como atrasos críticos "
     "e custos fora do padrão. Os dados utilizados são sintéticos e gerados de forma determinística, o que "
     "permite reproduzir exatamente os mesmos resultados a cada execução.")

titulo("1.2 Problema", 2)
para("Sem uma visão consolidada, a gestão da transportadora decide com base em percepções isoladas. Faltam "
     "respostas objetivas para perguntas como: qual percentual das entregas chega no prazo? Qual é o atraso "
     "médio quando ocorre? Em que região a operação é mais lenta ou mais cara? Quais motoristas são mais "
     "pontuais? Existem entregas com atrasos graves ou custos atípicos que merecem investigação? Além disso, os "
     "dados brutos precisam chegar à gestão em formatos adequados a cada uso: JSON para integração com outros "
     "sistemas, CSV para planilhas e HTML para leitura direta.")

titulo("1.3 Objetivos", 2)
para("**Objetivo geral:** construir um sistema organizado em microsserviços, que se comunicam por HTTP, capaz de "
     "armazenar os dados operacionais de entregas, calcular indicadores de desempenho e gerar um relatório final "
     "consolidado, com todos os serviços containerizados com Docker e orquestrados com Docker Compose.")
para("**Objetivos específicos:**", recuo=False)
itens([
    "implementar um Serviço de Dados com banco relacional (PostgreSQL) e endpoints de leitura e escrita;",
    "implementar um Serviço de Processamento que consome os dados exclusivamente via HTTP e aplica a lógica de "
    "cálculo de indicadores (pontualidade, atrasos, custo por quilômetro, ranking e outliers);",
    "implementar um Serviço de Relatório que orquestra os demais serviços e entrega o resultado em JSON, CSV e HTML;",
    "empacotar cada serviço em seu próprio Dockerfile e orquestrá-los com Docker Compose, com verificações de "
    "saúde (healthchecks) e ordem de inicialização;",
    "validar o funcionamento com testes automatizados e com um roteiro de verificação de ponta a ponta, "
    "registrando as evidências.",
])

titulo("1.4 Atendimento aos requisitos da atividade", 2)
para("A Tabela 1 relaciona cada requisito do enunciado da atividade à parte do sistema que o atende.")
tabela("Requisitos da atividade e como foram atendidos",
       ["Requisito", "Como foi atendido"],
       [["Pelo menos três microsserviços comunicando-se via HTTP",
         "Três serviços (Dados, Processamento e Relatório) que se chamam por HTTP/JSON, mais o banco PostgreSQL."],
        ["Serviço de Dados (banco + leitura/escrita)",
         "`servico-dados`, com PostgreSQL e endpoints GET, POST e PATCH sobre motoristas, veículos e entregas."],
        ["Serviço de Processamento",
         "`servico-processamento`, que obtém os dados via HTTP e calcula indicadores, ranking e outliers."],
        ["Serviço de Relatório (orquestração)",
         "`servico-relatorio`, que chama Dados e Processamento e gera o relatório em JSON, CSV ou HTML."],
        ["Um Dockerfile por serviço", "Três Dockerfiles, um em cada diretório de serviço."],
        ["Orquestração com Docker Compose",
         "`docker-compose.yml` com rede própria, healthchecks, `depends_on` e portas publicadas."]],
       [5.5, 10.5])

# ----------------------------------------------------------------------------
# 2 METODOLOGIA
# ----------------------------------------------------------------------------
titulo("2 Metodologia", 1)
titulo("2.1 Arquitetura geral", 2)
para("A arquitetura segue a ideia central de microsserviços descrita por Fowler e Lewis (2014): o sistema é "
     "dividido em serviços pequenos, cada um com uma responsabilidade bem definida, executando em seu próprio "
     "processo e comunicando-se por mecanismos leves, tipicamente uma API HTTP. Newman (2021) destaca que essa "
     "separação favorece a implantação independente e a evolução isolada de cada parte, ao custo de maior "
     "complexidade de comunicação e operação, aspecto tratado neste trabalho com healthchecks e tratamento de "
     "falhas entre serviços.")
para("A Figura 1 apresenta o diagrama de arquitetura. O cliente (navegador ou `curl`) acessa apenas o Serviço de "
     "Relatório (porta 8003). Este chama o Serviço de Processamento (porta 8002), que por sua vez consulta o "
     "Serviço de Dados (porta 8001). O Serviço de Relatório também consulta diretamente o Serviço de Dados para "
     "obter o inventário da base. Somente o Serviço de Dados acessa o PostgreSQL, que não tem porta publicada no "
     "computador hospedeiro.")
figura(DOCS / "diagrama_arquitetura.png", 15.0, "Arquitetura do sistema de microsserviços")

titulo("2.2 Tecnologias e justificativas", 2)
para("A Tabela 2 lista as tecnologias adotadas e o motivo de cada escolha. O critério geral foi usar o menor "
     "conjunto de ferramentas capaz de cumprir os requisitos, todas amplamente documentadas e fáceis de "
     "explicar e reproduzir.")
tabela("Tecnologias adotadas e justificativas",
       ["Tecnologia", "Papel", "Justificativa"],
       [["Python 3.12 e Flask 3.0.3", "API HTTP dos três serviços",
         "Microframework enxuto, com poucas linhas por endpoint e documentação extensa (PALLETS, s.d.; "
         "PYTHON SOFTWARE FOUNDATION, s.d.). Mantém o foco na arquitetura, não no framework."],
        ["Gunicorn 22.0.0", "Servidor WSGI nos contêineres",
         "Evita executar o servidor de desenvolvimento do Flask; dois processos de trabalho por serviço."],
        ["PostgreSQL 16.10", "Banco relacional do Serviço de Dados",
         "As entidades têm relacionamentos claros; chaves estrangeiras e restrições `CHECK` garantem "
         "integridade (POSTGRESQL GLOBAL DEVELOPMENT GROUP, s.d.)."],
        ["psycopg2-binary 2.9.9", "Acesso ao banco",
         "SQL puro e parametrizado (sem concatenação de texto), com pool de conexões; sem ORM, para que cada "
         "consulta seja explicável."],
        ["requests 2.32.3", "Cliente HTTP entre serviços",
         "Sessões, tempo limite e novas tentativas simples para chamadas entre Relatório, Processamento e Dados."],
        ["Docker", "Empacotamento de cada serviço",
         "Um Dockerfile por serviço, imagem base `python:3.12.12-slim` (versão de patch fixada), usuário sem "
         "privilégios; garante "
         "ambiente idêntico em qualquer máquina (DOCKER INC., s.d.)."],
        ["Docker Compose", "Orquestração local",
         "Sobe todos os serviços com um comando, define rede, ordem de inicialização e healthchecks "
         "(COMPOSE SPECIFICATION, s.d.)."],
        ["pytest 8.2.2", "Testes automatizados",
         "Valida cálculos, cliente HTTP, rotas e validações dos três serviços sem depender de Docker, rede "
         "ou banco (HTTP e banco simulados)."]],
       [3.6, 3.8, 8.6])

titulo("2.3 Serviço de Dados", 2)
para("O Serviço de Dados é o único componente que acessa o banco. Ele cria o esquema e carrega os dados "
     "sintéticos na inicialização, por meio do script `init_db.py`, executado antes do servidor Gunicorn. O "
     "script aguarda o banco ficar disponível (tentativas a cada 2 segundos, por até 60 segundos), aplica o "
     "esquema de forma idempotente e só semeia os dados se a tabela de motoristas estiver vazia, evitando "
     "duplicação quando o contêiner é reiniciado.")
para("A carga inicial usa uma semente fixa (42) e uma data base fixa, de modo que o conjunto é sempre idêntico: "
     "20 motoristas, 12 veículos e 400 entregas, com datas de saída entre março e maio de 2026. Cada motorista "
     "recebe uma propensão a atraso própria e cada região, um fator de distância e tráfego, para que a análise "
     "apresente diferenças reais entre eles. A Tabela 3 resume o modelo de dados.")
tabela("Modelo de dados do Serviço de Dados",
       ["Tabela", "Colunas principais", "Restrições"],
       [["motoristas", "id, nome, cnh_categoria, cidade_base, data_admissao", "Chave primária `id`"],
        ["veiculos", "id, placa, tipo, capacidade_kg, ano",
         "`placa` única; `capacidade_kg` > 0; tipos MOTO, VAN, CAMINHAO_LEVE e CAMINHAO"],
        ["entregas",
         "id, motorista_id, veiculo_id, regiao, status, distancia_km, custo, peso_kg, data_saida, "
         "prazo_previsto, entregue_em",
         "Chaves estrangeiras para motoristas e veículos; `status` em ENTREGUE, EM_ROTA, CANCELADA ou FALHA; "
         "`distancia_km` e `custo` não negativos; `peso_kg` > 0; `entregue_em` nulo quando não entregue"]],
       [2.6, 6.4, 7.0])
para("O atraso não é armazenado: é derivado de `entregue_em` e `prazo_previsto` pelo Serviço de Processamento, "
     "o que evita redundância e inconsistência. A Tabela 4 apresenta os endpoints. Entradas inválidas resultam "
     "em resposta JSON com status 400, recursos inexistentes em 404, método não permitido em 405, placa duplicada "
     "em 409, erro interno em 500 e banco indisponível em 503. Os timestamps são armazenados sem fuso horário: "
     "um deslocamento enviado na requisição (por exemplo `-03:00`) é descartado e o horário informado é mantido. "
     "Todas as consultas usam parâmetros do driver, nunca concatenação de texto.")
tabela("Endpoints do Serviço de Dados (porta 8001)",
       ["Método", "Caminho", "Descrição"],
       [["GET", "/health", "Saúde do serviço e do banco"],
        ["GET", "/resumo", "Contagem de motoristas, veículos e entregas"],
        ["GET", "/motoristas, /motoristas/<id>", "Lista paginada (filtro `cidade_base`) e detalhe"],
        ["POST", "/motoristas", "Cria motorista"],
        ["GET", "/veiculos, /veiculos/<id>", "Lista paginada (filtro `tipo`) e detalhe"],
        ["POST", "/veiculos", "Cria veículo (placa duplicada: 409)"],
        ["GET", "/entregas, /entregas/<id>",
         "Lista paginada com filtros `regiao`, `status`, `motorista_id`, `veiculo_id`, `data_inicio`, `data_fim`"],
        ["POST", "/entregas", "Cria entrega, validando chaves estrangeiras, status e datas"],
        ["PATCH", "/entregas/<id>", "Atualiza `status`, `entregue_em` e/ou `custo`"]],
       [2.0, 5.6, 8.4])
para("A paginação usa os parâmetros `limit` (padrão 100, máximo 1.000) e `offset`, e a resposta tem o formato "
     "`total`, `limit`, `offset` e `itens`.")

titulo("2.4 Serviço de Processamento", 2)
para("O Serviço de Processamento não possui driver de banco de dados: todos os dados chegam por HTTP do Serviço "
     "de Dados. O módulo `cliente_dados.py` percorre as páginas de 1.000 itens, usa tempo limite de 5 segundos e "
     "duas tentativas com espera de 1 segundo (pior caso de cerca de 11 segundos por chamada); se o Serviço de "
     "Dados estiver indisponível, a resposta é 502, e parâmetros inválidos resultam em 400. A lógica "
     "de negócio fica em `calculos.py`, composto por funções puras (sem entrada e saída), o que facilita o teste "
     "unitário. A Tabela 5 apresenta os indicadores e suas definições.")
tabela("Indicadores calculados pelo Serviço de Processamento",
       ["Indicador", "Definição"],
       [["Atraso (min)", "Para entregas ENTREGUE: máximo entre 0 e (`entregue_em` − `prazo_previsto`), em minutos"],
        ["No prazo", "Entrega concluída com atraso igual a 0"],
        ["% no prazo", "Entregas no prazo ÷ entregas concluídas × 100"],
        ["Atraso médio", "Média dos atrasos maiores que 0 (entregas atrasadas)"],
        ["Atraso médio geral", "Média do atraso sobre todas as entregas concluídas, incluindo as pontuais"],
        ["Custo por km", "Soma dos custos ÷ soma das distâncias, apenas entregas concluídas (0 se não houver distância)"],
        ["Ranking de motoristas",
         "Ordenado por % no prazo (decrescente), depois atraso médio (crescente) e nome; só entra quem tem pelo "
         "menos 5 entregas concluídas"],
        ["Atraso crítico", "Entrega concluída com atraso de 120 minutos ou mais"],
        ["Outlier de custo por km",
         "Regra do IQR sobre o custo por km de cada entrega: valores acima de Q3 + 1,5 × (Q3 − Q1)"]],
       [4.2, 11.8])
para("Os principais endpoints são `/metricas/resumo`, `/metricas/status`, `/metricas/regioes`, "
     "`/metricas/motoristas/ranking`, `/metricas/atrasos-criticos`, `/metricas/outliers-custo` e "
     "`/metricas/completo`. Este último reúne todos os blocos em uma única resposta, buscando os dados uma só "
     "vez, e é o utilizado pelo Serviço de Relatório. Todos aceitam os filtros opcionais `regiao`, "
     f"`data_inicio` e `data_fim`. As funções de cálculo são cobertas por testes automatizados "
     "(pontual, atrasada e cancelada, divisão por zero, desempate do ranking, exclusão por mínimo de entregas, "
     "limite de 119 contra 120 minutos e detecção de outliers); a suíte completa é descrita na Seção 2.8.")

titulo("2.5 Serviço de Relatório", 2)
para("O Serviço de Relatório orquestra o fluxo ao receber `GET /relatorio?formato=json|csv|html` (padrão "
     "`json`). Primeiro consulta `/resumo` no Serviço de Dados, depois `/metricas/completo` no Serviço de "
     "Processamento, repassando os filtros opcionais, e então monta um relatório consolidado com título, data "
     "e hora de geração, filtros, fontes consultadas, inventário da base, indicadores e observações "
     "automáticas (melhor e pior região, líder do ranking e contagem de casos críticos). A Tabela 6 descreve "
     "os formatos de saída.")
tabela("Formatos de saída do Serviço de Relatório",
       ["Formato", "Content-Type", "Descrição"],
       [["json", "application/json", "Documento consolidado com todas as seções"],
        ["csv", "text/csv; charset=utf-8",
         "Formato longo com colunas `secao,item,metrica,valor`, entregue como anexo `relatorio_entregas.csv`"],
        ["html", "text/html; charset=utf-8",
         "Página única com indicadores em destaque e tabelas de status, regiões, ranking, atrasos críticos e "
         "outliers; `lang=\"pt-BR\"` e cabeçalhos de tabela semânticos"]],
       [2.2, 4.6, 9.2])
para("Um valor de `formato` inválido resulta em 400. Se o Serviço de Dados ou o de Processamento falhar, "
     "expirar o tempo limite ou responder com erro, o Serviço de Relatório devolve 502 com JSON que identifica o "
     "serviço dependente que falhou, em vez de propagar uma exceção. Os tempos limite formam uma hierarquia: "
     "Processamento→Dados 5 segundos com duas tentativas (cerca de 11 segundos por chamada), "
     "Relatório→Processamento 30 segundos e `--timeout 60` do Gunicorn no Processamento e no Relatório. Assim "
     "cada camada espera mais do que a de baixo leva no pior caso, e a falha é reportada pelo serviço mais "
     "próximo da causa.")

titulo("2.6 Containerização e orquestração", 2)
para("Cada serviço possui seu próprio Dockerfile, baseado em `python:3.12.12-slim`. As dependências são instaladas "
     "antes da cópia do código, para aproveitar o cache de camadas, e o contêiner executa com um usuário sem "
     "privilégios. Todas as dependências estão fixadas em versões exatas. O Dockerfile do Serviço de Dados é "
     "mostrado a seguir; os demais seguem o mesmo padrão e diferem na porta e no comando de inicialização.")
codigo(DOCKERFILE)
para("O arquivo `docker-compose.yml` define o projeto `logistica` e a rede `logistica-net`, o volume `pgdata` do "
     "banco e quatro serviços: `db` (PostgreSQL 16.10), `servico-dados`, `servico-processamento` e "
     "`servico-relatorio`. Os pontos principais do desenho são:")
itens([
    "**Comunicação por nome de serviço:** as URLs internas usam os nomes do Compose (por exemplo "
    "`http://servico-dados:8001`), nunca `localhost`;",
    "**Healthchecks:** o banco usa `pg_isready`; os serviços Python consultam seu próprio `/health` com "
    "`urllib`, pois a imagem `slim` não inclui `curl`;",
    "**Ordem de inicialização:** `depends_on` com `condition: service_healthy` faz o Dados esperar o banco, o "
    "Processamento esperar o Dados e o Relatório esperar os dois;",
    "**Exposição mínima:** apenas as portas 8001, 8002 e 8003 são publicadas; o PostgreSQL fica restrito à rede "
    "interna;",
    "**Reinício automático:** `restart: unless-stopped` em todos os serviços;",
    "**Configuração por variáveis de ambiente**, com valores padrão e possibilidade de sobrescrita por arquivo "
    "`.env` (modelo em `.env.example`).",
])

titulo("2.7 Execução e verificação", 2)
para("Para executar o sistema são necessários apenas Docker e Docker Compose. Na raiz do projeto:")
codigo("docker compose up -d --build --wait\n"
       "docker compose ps\n"
       "curl http://localhost:8001/resumo\n"
       "curl http://localhost:8002/metricas/completo\n"
       "curl \"http://localhost:8003/relatorio?formato=html\"\n"
       "docker compose down -v")
para("A verificação consistiu em: (i) testes automatizados dos três serviços (Seção 2.8); (ii) validação do arquivo do Compose e "
     "subida com espera de saúde; (iii) consultas aos três serviços, incluindo filtros, criação de registros, "
     "duplicidade de placa (409) e recurso inexistente (404); (iv) os três formatos do relatório e o formato "
     "inválido (400); (v) teste de resiliência, parando o Serviço de Processamento e, depois, o de Dados; (vi) teste "
     "de idempotência, reiniciando o Serviço de Dados; (vii) reinício do banco, para observar a resposta do "
     "Serviço de Dados durante a queda; e (viii) recriação completa (`down -v` seguido de nova subida) para "
     "confirmar que os números são idênticos. As saídas foram capturadas na pasta `docs/evidencias`.")

titulo("2.8 Qualidade e testes", 2)
para(f"A suíte automatizada totaliza **{N_TESTES} testes** com pytest, executados com um único comando "
     f"(`python rodar_testes.py`), que roda o pytest uma vez por serviço: {T_DADOS} no Serviço de Dados, "
     f"{T_PROC} no de Processamento e {T_REL} no de Relatório. Os testes não usam Docker, rede nem banco: as "
     "chamadas HTTP e o acesso ao PostgreSQL são simulados. Eles cobrem os cálculos, o cliente HTTP (paginação, "
     "retentativas e tratamento de erro 400 e 5xx), as rotas dos três serviços, a validação de entrada do "
     "Serviço de Dados (inclusive a rejeição de NaN e infinito), o gerador de dados sintéticos e a conversão para "
     "CSV. As dependências de teste ficam em arquivos `requirements-dev.txt` e não entram nas imagens de "
     "produção. O código também foi verificado com o analisador estático Ruff, sem pendências.")
para("Alguns pontos de robustez foram tratados a partir de uma revisão de qualidade do código: hierarquia de "
     "tempos limite entre os serviços; tratamento de conexão perdida com o banco como indisponibilidade (503), "
     "sem mascarar o erro original; ausência de senha padrão no código do Serviço de Dados, que passa a falhar "
     "cedo se `DB_PASSWORD` não for definida; restrição `CHECK` para `peso_kg` positivo; e versões de patch "
     "fixadas para as imagens base.")

# ----------------------------------------------------------------------------
# 3 RESULTADOS
# ----------------------------------------------------------------------------
titulo("3 Resultados", 1)
titulo("3.1 Ambiente e conjunto de dados", 2)
para(f"Os resultados apresentados a seguir foram obtidos com a pilha completa em execução e o banco recém-criado, "
     f"de modo que refletem exclusivamente os dados semeados. O Serviço de Dados reportou **{BASE['motoristas']} "
     f"motoristas**, **{BASE['veiculos']} veículos** e **{n(BASE['entregas'], 0)} entregas**. O relatório "
     f"analisado foi gerado em {dh(rel['gerado_em'])}. A Tabela 7 mostra o estado dos contêineres e a "
     f"Figura 2, o número de entregas por status. Os {N_TESTES} testes automatizados dos três serviços foram "
     f"aprovados ({T_DADOS} em Dados, {T_PROC} em Processamento e {T_REL} em Relatório).")
tabela("Estado dos contêineres (docker compose ps)",
       ["Contêiner", "Serviço", "Saúde", "Porta (host→contêiner)"],
       TAB_PS, [5.6, 4.0, 2.0, 4.4], ["l", "l", "c", "l"],
       fonte="Fonte: docs/evidencias/docker-compose-ps.txt.")
figura(DOCS / "grafico_status.png", 12.5, "Entregas por status",
       "Fonte: docs/evidencias/relatorio.json.")
para(f"Das {R['total']} entregas, {STATUS['ENTREGUE']} foram concluídas, {STATUS['EM_ROTA']} estavam em rota, "
     f"{STATUS['FALHA']} falharam e {STATUS['CANCELADA']} foram canceladas. Portanto, {nao_concluidas} entregas "
     f"({pct(round(nao_concluidas / R['total'] * 100, 2))}) não foram concluídas e ficam fora dos indicadores "
     "de prazo e custo, que consideram apenas entregas concluídas.")

titulo("3.2 Indicadores gerais", 2)
tabela("Indicadores gerais de desempenho",
       ["Indicador", "Valor"],
       [["Entregas analisadas", n(R["total"], 0)],
        ["Entregas concluídas", n(R["finalizadas"], 0)],
        ["Entregas no prazo", n(R["no_prazo"], 0)],
        ["Entregas atrasadas", n(R["atrasadas"], 0)],
        ["% no prazo (sobre concluídas)", pct(R["pct_no_prazo"])],
        ["Atraso médio das atrasadas", f"{n(R['atraso_medio_min'])} min"],
        ["Atraso médio geral (todas as concluídas)", f"{n(R['atraso_medio_geral_min'])} min"],
        ["Atraso máximo", f"{n(R['atraso_maximo_min'], 0)} min"],
        ["Distância total (concluídas)", f"{n(R['distancia_total_km'])} km"],
        ["Custo total (concluídas)", f"R$ {n(R['custo_total'])}"],
        ["Custo por km", f"R$ {n(R['custo_por_km'])}"]],
       [10.0, 6.0], ["l", "r"], fonte="Fonte: docs/evidencias/relatorio.json.")
para(f"Das {R['finalizadas']} entregas concluídas, {R['no_prazo']} chegaram no prazo, o que corresponde a "
     f"{pct(R['pct_no_prazo'])}; as outras {R['atrasadas']} ({pct(pct_atrasadas)}) atrasaram. Quando há atraso, "
     f"ele é relevante: {n(R['atraso_medio_min'])} minutos em média, com máximo de {R['atraso_maximo_min']} "
     f"minutos. Como a maior parte das entregas é pontual, o atraso médio geral cai para "
     f"{n(R['atraso_medio_geral_min'])} minutos. Esse contraste mostra por que a média das atrasadas e a média "
     "geral devem ser lidas juntas: a primeira descreve a gravidade do problema, a segunda, o seu impacto no "
     f"conjunto. O custo médio da operação foi de R$ {n(R['custo_por_km'])} por quilômetro.")

titulo("3.3 Desempenho por região", 2)
tabela("Pontualidade e custo por região",
       ["Região", "Total", "Concluídas", "% no prazo", "Atraso médio (min)", "Custo/km (R$)"],
       [[r["regiao"], r["total"], r["entregues"], n(r["pct_no_prazo"]), n(r["atraso_medio_min"]),
         n(r["custo_por_km"])] for r in REGIOES],
       [3.2, 1.8, 2.6, 2.6, 3.2, 2.6], ["l", "r", "r", "r", "r", "r"],
       fonte="Fonte: docs/evidencias/relatorio.json.")
figura(DOCS / "grafico_regioes.png", 12.5, "Percentual de entregas no prazo por região",
       "Fonte: docs/evidencias/relatorio.json.")
texto_regiao = (f"A região mais pontual é {melhor['regiao']}, com {pct(melhor['pct_no_prazo'])} das entregas no "
                f"prazo, e a menos pontual é {pior['regiao']}, com {pct(pior['pct_no_prazo'])}: uma diferença de "
                f"{n(dif_regioes)} pontos percentuais. ")
if menor_atraso["regiao"] != melhor["regiao"]:
    texto_regiao += (f"Pontualidade e gravidade do atraso não andam juntas: {melhor['regiao']} tem o melhor "
                     f"percentual, mas, entre as entregas atrasadas, o atraso médio é de "
                     f"{n(melhor['atraso_medio_min'])} min, contra {n(menor_atraso['atraso_medio_min'])} min em "
                     f"{menor_atraso['regiao']}, que tem o menor atraso médio. A região com maior atraso médio é "
                     f"{maior_atraso['regiao']} ({n(maior_atraso['atraso_medio_min'])} min). ")
texto_regiao += (f"Em custo, {mais_barata['regiao']} é a mais econômica (R$ {n(mais_barata['custo_por_km'])} por "
                 f"km) e {mais_cara['regiao']} a mais cara (R$ {n(mais_cara['custo_por_km'])} por km). "
                 "Como os dados são sintéticos e cada região recebeu parâmetros próprios de distância e tráfego, "
                 "essas diferenças demonstram a capacidade do sistema de evidenciá-las, não conclusões sobre uma "
                 "cidade real.")
para(texto_regiao)

titulo("3.4 Desempenho dos motoristas", 2)
para("O ranking considera apenas motoristas com pelo menos 5 entregas concluídas. A Tabela 10 apresenta as dez "
     "primeiras posições e a Figura 4, o percentual de pontualidade correspondente.")
tabela("Ranking dos 10 motoristas mais pontuais",
       ["Pos.", "Motorista", "Concluídas", "% no prazo", "Atraso médio (min)", "Custo/km (R$)"],
       [[r["posicao"], r["nome"], r["entregas_finalizadas"], n(r["pct_no_prazo"]), n(r["atraso_medio_min"]),
         n(r["custo_por_km"])] for r in RANK],
       [1.2, 4.8, 2.6, 2.2, 3.0, 2.2], ["c", "l", "r", "r", "r", "r"],
       fonte="Fonte: docs/evidencias/relatorio.json.")
figura(DOCS / "grafico_ranking.png", 13.0, "Dez motoristas com maior percentual de entregas no prazo",
       "Fonte: docs/evidencias/relatorio.json.")
para(f"O líder é {lider['nome']}, com {pct(lider['pct_no_prazo'])} de pontualidade em "
     f"{lider['entregas_finalizadas']} entregas concluídas; a décima posição ({ultimo['nome']}) tem "
     f"{pct(ultimo['pct_no_prazo'])}. O ranking também revela um caso que o percentual isolado esconde: "
     f"{maior_atraso_motorista['nome']}, em {maior_atraso_motorista['posicao']}º lugar com "
     f"{pct(maior_atraso_motorista['pct_no_prazo'])} no prazo, tem o maior atraso médio entre os dez "
     f"({n(maior_atraso_motorista['atraso_medio_min'])} min), ou seja, atrasa pouco, mas, quando atrasa, o "
     "atraso é grande. Por isso o critério de desempate e a coluna de atraso médio fazem parte da análise.")

titulo("3.5 Atrasos críticos e custos atípicos", 2)
para(f"Foram identificadas {CRIT['total']} entregas com atraso crítico (120 minutos ou mais), listadas na "
     f"Tabela 11. O maior atraso, de {crit_max['atraso_min']} minutos, ocorreu na entrega {crit_max['id']} "
     f"({crit_max['regiao']}).")
tabela(f"Entregas com atraso crítico (limite de {CRIT['limite_min']} min)",
       ["Entrega", "Motorista (id)", "Região", "Atraso (min)", "Prazo previsto", "Entregue em"],
       [[c["id"], c["motorista_id"], c["regiao"], c["atraso_min"], dh(c["prazo_previsto"]), dh(c["entregue_em"])]
        for c in CRIT["itens"]],
       [1.8, 2.4, 2.6, 2.2, 3.5, 3.5], ["r", "r", "l", "r", "c", "c"],
       fonte="Fonte: docs/evidencias/relatorio.json.")
ids_out = ", ".join(str(i) for i in OUT["ids"])
custos_out = "; ".join(f"entrega {i['id']}: R$ {n(i['custo_km'])}/km" for i in OUT["itens"])
para(f"Na análise de custo por quilômetro, a regra do IQR considerou {OUT['analisadas']} entregas concluídas, com "
     f"Q1 = {n(OUT['q1'])}, Q3 = {n(OUT['q3'])} e IQR = {n(OUT['iqr'])}, o que resulta em limite superior de "
     f"R$ {n(OUT['limite_superior'])} por km. {OUT['total']} entregas ultrapassaram esse limite (ids {ids_out}): "
     f"{custos_out}. Esses registros são candidatos à auditoria, por exemplo, para verificar erro de "
     "lançamento ou rota anormal.")

titulo("3.6 Saídas do Serviço de Relatório", 2)
obs = rel["observacoes"]
para("O endpoint `/relatorio` respondeu com status 200 nos três formatos (`application/json`, "
     "`text/csv; charset=utf-8` e `text/html; charset=utf-8`) e com status 400 para um formato inválido. As "
     "observações automáticas incluídas no relatório consolidado foram: " + " ".join(obs) + " A Figura 5 "
     "mostra a parte inicial da página HTML gerada, e o formato CSV (`docs/evidencias/relatorio.csv`) "
     "reúne os mesmos números em formato longo, adequado para planilhas.")
if SCREENSHOT.exists():
    figura(SCREENSHOT, 9.5, "Relatório em HTML gerado pelo Serviço de Relatório (captura da página)",
           "Fonte: docs/evidencias/relatorio-html.png.")

titulo("3.7 Resiliência e idempotência", 2)
para("Dois comportamentos foram testados. No teste de resiliência, o Serviço de Processamento foi parado e o "
     "relatório foi solicitado: a resposta foi HTTP 502, com JSON contendo a mensagem de falha e o nome do "
     "serviço dependente (`processamento`). Após iniciar novamente o serviço, o relatório voltou a responder "
     "com HTTP 200. Repetido com o Serviço de Dados parado, o relatório respondeu 502 nomeando o serviço `dados`, "
     "e o Serviço de Processamento, consultado diretamente, respondeu 502 em poucos segundos, antes do tempo "
     "limite do Relatório. Isso mostra que a falha de um serviço é isolada e comunicada de forma controlada, "
     "sem derrubar o Serviço de Relatório.")
para("No reinício do banco, o Serviço de Dados foi consultado uma vez por segundo durante 40 segundos: houve "
     "respostas 503 (banco indisponível) apenas no instante da queda, nenhuma resposta 500, e as demais foram "
     "200, ou seja, o serviço se recuperou sozinho. Os números exatos dessa execução estão em "
     "`docs/evidencias/resiliencia.txt`.")
para("No teste de idempotência, o Serviço de Dados foi reiniciado. O resumo permaneceu em 400 entregas, 20 "
     "motoristas e 12 veículos, e o registro do contêiner indicou que o banco já estava populado e que a "
     "semeadura foi ignorada. Por fim, a recriação completa do ambiente (`down -v` e nova subida) reproduziu "
     "exatamente os mesmos números, confirmando o determinismo da carga de dados.")

titulo("3.8 Análise crítica e limitações", 2)
para("Os resultados mostram que a divisão em três serviços funcionou como pretendido: o Serviço de Dados "
     "concentra o armazenamento, o de Processamento concentra as regras de cálculo, que podem ser testadas "
     "sem rede ou banco, e o de Relatório apenas orquestra e formata. A mudança de formato de saída, por "
     "exemplo, não exige tocar nas regras de negócio. Em contrapartida, a comunicação síncrona cria "
     "dependência em cadeia, conforme discutido por Newman (2021), e o teste de resiliência evidencia que a "
     "queda do Processamento impede o relatório, ainda que de forma controlada. As limitações a considerar são:")
itens([
    "**Sem autenticação nem HTTPS:** todos os endpoints, inclusive os de escrita (POST e PATCH), são abertos; "
    "o uso é estritamente didático;",
    "**Dados sintéticos:** os padrões regionais e de motoristas foram parametrizados na geração, portanto não "
    "representam uma operação real;",
    "**Sem cache, fila ou escalabilidade horizontal:** cada relatório refaz toda a consulta e o cálculo, e a "
    "comunicação é síncrona;",
    "**Senha didática do banco** presente apenas no arquivo do Compose (o código do Serviço de Dados não "
    "tem senha padrão), aceitável somente para fins didáticos;",
    "**Duplicação deliberada:** pequenas utilidades (leitura dos filtros, conversão de parâmetros inteiros e "
    "tratamento de erros HTTP) se repetem em mais de um serviço para mantê-los independentes, sem biblioteca "
    "compartilhada;",
    "**Dependências sem lockfile completo:** Flask, Werkzeug e Jinja2 estão fixados, mas as demais dependências "
    "transitivas não, e as imagens base são fixadas por versão de patch, não por digest;",
    "**Consistência da leitura:** o Processamento lê os dados em páginas, sem transação única, de modo que "
    "gravações concorrentes durante o cálculo poderiam gerar uma visão ligeiramente inconsistente;",
    "**Verificação restrita ao ambiente local** com Docker Compose; não houve teste de carga nem implantação em "
    "nuvem.",
])

# ----------------------------------------------------------------------------
# 4 CONCLUSÃO
# ----------------------------------------------------------------------------
titulo("4 Conclusão", 1)
para("A atividade resultou em um sistema de três microsserviços que se comunicam por HTTP, apoiados em um banco "
     "PostgreSQL, todos containerizados com um Dockerfile por serviço e orquestrados com Docker Compose. Os "
     "objetivos propostos foram cumpridos: o Serviço de Dados oferece leitura e escrita sobre um modelo "
     "relacional; o Serviço de Processamento consome esses dados exclusivamente por HTTP e calcula os "
     "indicadores; e o Serviço de Relatório orquestra os dois e entrega o resultado em JSON, CSV e HTML.")
para(f"Do ponto de vista dos resultados, o sistema permitiu responder às perguntas que motivaram o trabalho: "
     f"{pct(R['pct_no_prazo'])} das {R['finalizadas']} entregas concluídas chegaram no prazo, {melhor['regiao']} "
     f"é a região mais pontual e {pior['regiao']}, a menos pontual, {lider['nome']} lidera o ranking de "
     f"motoristas, e foram sinalizados {CRIT['total']} atrasos críticos e {OUT['total']} custos atípicos. Os "
     "testes de resiliência e idempotência e a reprodução exata dos números após a recriação do ambiente "
     "reforçam a confiabilidade do desenho.")
para("O principal aprendizado foi perceber, na prática, os benefícios e os custos dos microsserviços: a "
     "separação de responsabilidades facilita testar, entender e evoluir cada parte, mas exige cuidado com a "
     "ordem de inicialização, a verificação de saúde, o tratamento de falhas entre serviços e a configuração "
     "de rede. O uso de Docker e Compose tornou o ambiente totalmente reproduzível com um único comando.")
para("Como trabalhos futuros, propõem-se: (i) autenticação e autorização, com HTTPS; (ii) mensageria ou fila "
     "para desacoplar o cálculo da geração do relatório; (iii) cache dos indicadores; (iv) observabilidade, com "
     "logs estruturados, métricas e rastreamento distribuído; (v) integração contínua com execução automática "
     "dos testes e da construção das imagens; e (vi) testes de integração e de carga automatizados.")

# ----------------------------------------------------------------------------
# REFERÊNCIAS
# ----------------------------------------------------------------------------
titulo("Referências", 1)
ACESSO = "Acesso em: [[DATA DE ACESSO]]."
refs = [
    "COMPOSE SPECIFICATION. **The Compose Specification**. [s.d.]. Disponível em: https://compose-spec.io/. "
    + ACESSO,
    "DOCKER INC. **Docker Docs**: Docker Engine e Docker Compose. [s.d.]. Disponível em: "
    "https://docs.docker.com/. " + ACESSO,
    "FOWLER, Martin; LEWIS, James. **Microservices**: a definition of this new architectural term. 2014. "
    "Disponível em: https://martinfowler.com/articles/microservices.html. " + ACESSO,
    "NEWMAN, Sam. **Building Microservices**: designing fine-grained systems. 2. ed. Sebastopol: O'Reilly "
    "Media, 2021.",
    "PALLETS. **Flask Documentation**. [s.d.]. Disponível em: https://flask.palletsprojects.com/. " + ACESSO,
    "POSTGRESQL GLOBAL DEVELOPMENT GROUP. **PostgreSQL Documentation**. [s.d.]. Disponível em: "
    "https://www.postgresql.org/docs/. " + ACESSO,
    "PYTHON SOFTWARE FOUNDATION. **Python 3 Documentation**. [s.d.]. Disponível em: "
    "https://docs.python.org/3/. " + ACESSO,
    "PYTHON-DOCX. **python-docx Documentation**. [s.d.]. Disponível em: "
    "https://python-docx.readthedocs.io/. " + ACESSO,
    "PYTHON-PPTX. **python-pptx Documentation**. [s.d.]. Disponível em: "
    "https://python-pptx.readthedocs.io/. " + ACESSO,
]
for ref in refs:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(12)
    runs(p, ref)

doc.core_properties.title = "Sistema de Microsserviços para Gestão e Análise de Desempenho de Entregas em Logística Urbana"
doc.core_properties.author = ""
doc.save(SAIDA)
print(f"Gerado: {SAIDA}")


